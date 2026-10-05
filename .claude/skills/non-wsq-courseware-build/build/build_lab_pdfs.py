#!/usr/bin/env python3
"""Render every C1123 lab's instructions to labs/lab-NN-<slug>/LAB-NN-Instructions.pdf.

The content is the same block stream build_lab_guides.py writes to
LAB-NN-Instructions.md (lab_blocks), so the Markdown and PDF never diverge.
The intermediate DOCX is built in a temporary folder; nothing in the repo is deleted.
Run build_lab_guides.py first so assets/expected-evidence.png exists.
"""
import os, shutil, subprocess, sys, tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
import prodoc
from build_lab_guides import ACTS, LABS, folder_name, lab_blocks

BRAND = RGBColor(0x1F, 0x6F, 0xEB)
INKCODE = RGBColor(0x0B, 0x30, 0x60)
GREY = RGBColor(0x55, 0x5B, 0x66)
LOGO = os.path.join(os.path.dirname(HERE), "assets", "tertiary-infotech-logo.png")


def shade(par, fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def code(doc, text):
    par = doc.add_paragraph(); shade(par, "EAF2FF")
    par.paragraph_format.left_indent = Pt(10); par.paragraph_format.space_after = Pt(6)
    r = par.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(9.5); r.font.color.rgb = INKCODE


def render(a, folder, out_docx):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10.5)
    prodoc.style_headings(doc)
    prodoc.add_cover_page(doc, f"LAB {a['num']:02d} INSTRUCTIONS", a["title"], C.VERSION.lstrip("v"),
                          org_logo=LOGO, course_logo=None, course_code=C.COURSE_CODE)
    for blk in lab_blocks(a):
        k = blk[0]
        if k == "h1":
            doc.add_heading(blk[1], level=1)
        elif k == "meta":
            for label, value in blk[1]:
                par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(1)
                r = par.add_run(f"{label}: "); r.bold = True; r.font.color.rgb = BRAND
                par.add_run(value)
        elif k == "h2":
            h = doc.add_heading(blk[1], level=2); h.paragraph_format.keep_with_next = True
        elif k == "p":
            doc.add_paragraph(blk[1])
        elif k == "bullets":
            for x in blk[1]:
                doc.add_paragraph(x, style="List Bullet")
        elif k == "numbered":
            for i, x in enumerate(blk[1], 1):
                par = doc.add_paragraph(); par.paragraph_format.left_indent = Pt(14)
                r = par.add_run(f"{i}.  "); r.bold = True; r.font.color.rgb = BRAND
                par.add_run(x)
        elif k == "steps":
            for i, (text, cmd) in enumerate(blk[1], 1):
                par = doc.add_paragraph(); par.paragraph_format.left_indent = Pt(14)
                par.paragraph_format.keep_with_next = bool(cmd)
                r = par.add_run(f"{i}.  "); r.bold = True; r.font.color.rgb = BRAND
                par.add_run(text)
                if cmd:
                    code(doc, cmd)
        elif k == "code":
            code(doc, blk[1])
        elif k == "table":
            t = doc.add_table(rows=1, cols=len(blk[1])); t.style = "Table Grid"
            for cell, head in zip(t.rows[0].cells, blk[1]):
                cell.text = ""; r = cell.paragraphs[0].add_run(head); r.bold = True; r.font.size = Pt(9.5)
                prodoc._shade_cell(cell, "EAF2FF")
            for path, desc in blk[2]:
                cells = t.add_row().cells
                for cell, val, mono in ((cells[0], path, True), (cells[1], desc, False)):
                    cell.text = ""; r = cell.paragraphs[0].add_run(val); r.font.size = Pt(9.5)
                    if mono: r.font.name = "Consolas"
        elif k == "img":
            p = os.path.join(folder, blk[1])
            if os.path.exists(p):
                doc.add_picture(p, width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(blk[2]); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY
        elif k == "note":
            par = doc.add_paragraph(); shade(par, "FFF4E5")
            r = par.add_run("Note: "); r.bold = True
            par.add_run(blk[1])
    prodoc.add_page_numbers(doc, left_text=f"{C.TITLE} · {C.COURSE_CODE} · Lab {a['num']:02d}")
    doc.save(out_docx)


def main():
    only = {int(x) for x in sys.argv[1:] if x.isdigit()}
    with tempfile.TemporaryDirectory() as tmp:
        for a in ACTS:
            if only and a["num"] not in only:
                continue
            folder = os.path.join(LABS, folder_name(a))
            stem = f"LAB-{a['num']:02d}-Instructions"
            docx_path = os.path.join(tmp, stem + ".docx")
            render(a, folder, docx_path)
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, docx_path],
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            shutil.copyfile(os.path.join(tmp, stem + ".pdf"), os.path.join(folder, stem + ".pdf"))
            print("Saved", os.path.join(folder_name(a), stem + ".pdf"))


if __name__ == "__main__":
    main()
