"""Crop the reusable diagrams/screenshots out of the private v3 reference deck.

For each selected slide the crop box is the union of its picture/group/diagram
shapes plus any annotation text boxes that overlap them, so the old slide title,
footer and bullet text are left behind. Output: courseware/assets/reference-diagrams/ref-NNN.png
"""
from pathlib import Path
import subprocess, sys, tempfile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parent.parent
REF = next((ROOT/'reference').glob('*.pptx'))
OUT = ROOT/'courseware/assets/reference-diagrams'
# Slide numbers used by course_data.TOPIC_SLIDES
ONLY = [int(a) for a in sys.argv[1:]]
SLIDES = [20,21,22,23,24,28,29,30,31,38,39,42,43,44,45,46,48,55,57,58,65,68,69,70,71,72,79,
          86,88,112,115,116,117,137,138,139,140,141,144,147,150,162,181,182,189,209,214,
          216,220,223,240,248,251,253,254,279,283,289,291,54]
DPI = 200
FIX = {65:(0.0,0.27,1,1), 214:(0,0.07,1,1), 216:(0,0.06,1,0.86), 150:(0,0,1,0.80)}

def visual(sh):
    t = sh.shape_type
    if t in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.CHART,
             MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM):
        return True
    if t == MSO_SHAPE_TYPE.AUTO_SHAPE and not (sh.has_text_frame and len(sh.text_frame.text) > 40):
        return True
    return False

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(REF); W, H = prs.slide_width, prs.slide_height
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(tmp),str(REF)],check=True,capture_output=True)
    pdf = next(tmp.glob('*.pdf'))
    slides = list(prs.slides)
    for n in (ONLY or SLIDES):
        s = slides[n-1]
        boxes = [(sh.left, sh.top, sh.left+sh.width, sh.top+sh.height) for sh in s.shapes
                 if sh.width and sh.height and visual(sh) and sh.top > H*0.12 and sh.top < H*0.93]
        if not boxes:   # text-built diagram: take the whole body below the title
            boxes=[(0,int(H*0.15),W,int(H*0.93))]
        x0=min(b[0] for b in boxes); y0=min(b[1] for b in boxes); x1=max(b[2] for b in boxes); y1=max(b[3] for b in boxes)
        # include short annotation text boxes that overlap the visual area
        for sh in s.shapes:
            if not sh.has_text_frame or visual(sh) or not sh.width: continue
            bx=(sh.left, sh.top, sh.left+sh.width, sh.top+sh.height)
            if bx[1] < H*0.12 or bx[1] > H*0.93: continue
            if bx[0] < x1 and bx[2] > x0 and bx[1] < y1 and bx[3] > y0 and len(sh.text_frame.text) < 160:
                x0=min(x0,bx[0]); y0=min(y0,bx[1]); x1=max(x1,bx[2]); y1=max(y1,bx[3])
        x0=max(0,x0); y0=max(int(H*0.11),y0); x1=min(W,x1); y1=min(int(H*0.94),y1)
        png = tmp/f'p{n}'
        subprocess.run(['pdftoppm','-r',str(DPI),'-png','-singlefile','-f',str(n),'-l',str(n),str(pdf),str(png)],check=True)
        im = Image.open(str(png)+'.png').convert('RGB')
        sx, sy = im.width/W, im.height/H
        crop = im.crop((int(x0*sx), int(y0*sy), int(x1*sx), int(y1*sy)))
        # trim surrounding white
        bg = Image.new('RGB', crop.size, (255,255,255))
        bb = ImageChops.difference(crop, bg).getbbox()
        if bb: crop = crop.crop((max(0,bb[0]-8),max(0,bb[1]-8),min(crop.width,bb[2]+8),min(crop.height,bb[3]+8)))
        fx = FIX.get(n)
        if fx:  # trim clipped text fragments: fractions (left, top, right, bottom)
            crop = crop.crop((int(fx[0]*crop.width), int(fx[1]*crop.height), int(fx[2]*crop.width), int(fx[3]*crop.height)))
        crop.save(OUT/f'ref-{n:03d}.png')
        print(n, crop.size)

if __name__ == '__main__':
    main()
